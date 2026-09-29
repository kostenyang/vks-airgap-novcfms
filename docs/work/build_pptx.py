"""Build the VKS air-gap technical deck on the official Tech Tuesday vSAN 9.1 template (vcf-base).
New content slides are added from the template's own layouts (by NAME), cover + Thank You are the template's pages.
Then clone-slide.py keep reorders: cover, new slides..., Thank You."""
import copy, os, subprocess, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
DIA = os.path.join(HERE, "diagrams")
BASE = os.path.join(HERE, "base.pptx")
TMP = os.path.join(HERE, "deck-build.pptx")
OUT = os.path.join(os.path.dirname(HERE), "VKS-AirGap-VCF911-NoVCFMS.pptx")
HEX = lambda h: RGBColor.from_string(h.lstrip("#"))
PLUM, TEAL, SKY, GREEN, GOLD, LINE, NAVY = "#6C4B94", "#007B8C", "#0098C7", "#61A60E", "#F3BA16", "#717074", "#1B1D36"

prs = Presentation(BASE)
M1 = prs.slide_masters[0]
LAY = {l.name: l for l in M1.slide_layouts}


def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(idx)


def set_text(shape, text, size=None, bold=None, color=None):
    tf = shape.text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            if size: r.font.size = Pt(size)
            if bold is not None: r.font.bold = bold
            if color: r.font.color.rgb = HEX(color)


def bullets(shape, items, size=None):
    tf = shape.text_frame
    tf.clear()
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = it
        if size:
            for r in p.runs: r.font.size = Pt(size)


def drop(shape):
    shape._element.getparent().remove(shape._element)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new(layout, title=None, sub=None):
    s = prs.slides.add_slide(LAY[layout])
    if title is not None:
        ph(s, 0).text = title
    if sub is not None:
        ph(s, 10).text = sub
    return s


def section(layout, title, sub, note):
    s = new(layout, title, sub); notes(s, note); return s


def diagram_slide(title, sub, png, items, note):
    s = new("Diagram with Content on Right", title, sub)
    area = ph(s, 14); L, T, W, H = area.left, area.top, area.width, area.height; drop(area)
    s.shapes.add_picture(os.path.join(DIA, png), L, T, W, H)
    bullets(ph(s, 17), items, 13)
    notes(s, note); return s


def table(slide, rows, colw, left, top, width, font=11, header=PLUM, rowh=0.3, bold_first_col=False):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, left, top, width, Inches(rowh * nr))
    t = shp.table
    tot = sum(colw)
    for i, w in enumerate(colw):
        t.columns[i].width = int(width * w / tot)
    for r in range(nr):
        t.rows[r].height = Inches(rowh)
        for c in range(nc):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            tf.text = rows[r][c]
            for p in tf.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font + (0.5 if r == 0 else 0))
                    run.font.bold = (r == 0) or (bold_first_col and c == 0)
                    run.font.color.rgb = HEX("#FFFFFF" if r == 0 else "#333333")
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEX(header if r == 0 else ("#F2F2F2" if r % 2 == 0 else "#FFFFFF"))
    return shp


def table_slide(title, sub, rows, colw, note, font=11, header=PLUM, rowh=0.3, bold_first_col=True):
    s = new("Title, Subtitle and Content", title, sub)
    area = ph(s, 14); L, T, W = area.left, area.top, area.width; drop(area)
    table(s, rows, colw, L, T, W, font, header, rowh, bold_first_col)
    notes(s, note); return s


def text_slide(title, sub, items, note, size=18):
    s = new("Title, Subtitle and Content", title, sub)
    tf = ph(s, 14).text_frame; tf.clear()
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        m = p.add_run(); m.text = "▍ "; m.font.size = Pt(size); m.font.color.rgb = HEX([PLUM, TEAL, SKY, GREEN][i % 4])
        r = p.add_run(); r.text = it; r.font.size = Pt(size)
        p.space_after = Pt(14)
    notes(s, note); return s


def box(slide, x, y, w, h, head, lines, color):
    from pptx.enum.shapes import MSO_SHAPE
    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.45))
    hdr.fill.solid(); hdr.fill.fore_color.rgb = HEX(color); hdr.line.fill.background()
    set_text(hdr, head, 15, True, "#FFFFFF"); hdr.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.45), Inches(w), Inches(h - 0.45))
    body.fill.solid(); body.fill.fore_color.rgb = HEX("#FFFFFF"); body.line.color.rgb = HEX(color); body.line.width = Pt(1.5)
    tf = body.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = Inches(0.1)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + ln
        for r in p.runs: r.font.size = Pt(13); r.font.color.rgb = HEX("#333333")
        p.space_after = Pt(6)


def code_slide(title, sub, code, items, note):
    from pptx.enum.shapes import MSO_SHAPE
    s = new("Diagram with Content on Right", title, sub)
    area = ph(s, 14); L, T, W, H = area.left, area.top, area.width, area.height; drop(area)
    bx = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, T, W, H)
    bx.fill.solid(); bx.fill.fore_color.rgb = HEX("#F2F2F2"); bx.line.color.rgb = HEX(LINE); bx.line.width = Pt(1)
    tf = bx.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.15); tf.margin_top = Inches(0.1)
    for i, ln in enumerate(code.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln if ln else " "
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.name = "Consolas"; r.font.size = Pt(10.5)
            r.font.color.rgb = HEX(PLUM if ln.strip().startswith("#") else "#1B1D36")
    bullets(ph(s, 17), items, 13)
    notes(s, note); return s


new_start = len(prs.slides) + 1   # first new slide number (template has 30)

# ---------------- cover is template slide 1 (edited below); overview
s = new("Title Only", "Air-Gapped VKS without VCF Management Services — at a Glance")
box(s, 0.63, 1.3, 2.95, 3.3, "Platform", ["VCF bring-up 跑到 NSX 即中止", "3 × nested ESXi 9.1.1 + vSAN", "vCenter · SDDC Manager · NSX", "無 VSP / VCFMS / VCF Ops"], PLUM)
box(s, 3.73, 1.3, 2.95, 3.3, "Supervisor", ["VDS 網路 + Foundation LB (one-arm)", "3 CP + 2 FLB，API VIP 130.200", "ESXi = Spherelet agent node", "內建 VKS 服務 3.7.0"], TEAL)
box(s, 6.83, 1.3, 2.95, 3.3, "Air-Gap Supply Chain", ["bastion 下載 → admin 匯入", "Harbor v2.15.2 離線安裝", "VKr × 3 → 本地 content library", "VCF CLI + 13 plugins 離線安裝"], SKY)
box(s, 9.93, 1.3, 2.95, 3.3, "VKS & Add-ons", ["vks-test01（ClusterClass v3.7.0）", "addon 套件庫改指 Harbor（官方 override）", "DNS air-gap 下驗證通過", "0 個外部 registry image"], GREEN)
from pptx.enum.shapes import MSO_SHAPE
bn = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.63), Inches(4.95), Inches(12.25), Inches(1.1))
bn.fill.solid(); bn.fill.fore_color.rgb = HEX("#E8F0F4"); bn.line.fill.background()
set_text(bn, "關鍵發現：VKS 3.7 預設 addon 套件庫指向 projects.packages.broadcom.com — 看起來「Reconcile succeeded」其實在用外網。"
         "以官方 override 改指 Harbor，並在 DNS 切斷下完成反向對照驗證。", 15, False, NAVY)
bn.text_frame.word_wrap = True; bn.text_frame.margin_left = Inches(0.25)
notes(s, "總覽：四大塊。重點是「沒有 VCFMS 也做得到 air-gap VKS」，但 addon 套件庫必須用官方文件化的 override 改指內部 Harbor，並在切斷外網下驗證。")

# ---------------- section 1
section("Section - VKS", "Architecture & Design Decisions", "架構與設計決策",
        "第一章：先看整體架構，再看幾個決定性的設計選擇。")
diagram_slide("One Air-Gapped Supervisor Built from Standard Components", "整體架構：外層支援 VM + nested VCF 9.1.1 管理網域",
              "d1-architecture.png",
              ["外層：bastion（唯一可上網）、Harbor、admin、DNS/depot、pfSense",
               "nested：vCenter / SDDC Manager / NSX 由 bring-up 建立",
               "Supervisor 走 VDS + Foundation LB，不經 NSX",
               "管理網 120.x、workload 網 130.x（同一 L2、不同 subnet）",
               "air-gap 邊界：lab DNS 沒有 upstream"],
              "左邊是外層 vSphere 上的支援 VM，只有 bastion 在網際網路那一側。右邊是 nested 的 VCF 管理網域。"
              "Supervisor 用 VDS + FLB，所以 NSX 雖然裝了但不在 Supervisor 的資料路徑上。")
table_slide("Design Decisions That Shape an Air-Gapped VKS", "關鍵設計決策",
            [["決策", "選擇", "理由"],
             ["平台建置", "VCF bring-up 跑到 NSX 後中止", "spec 缺 vspClusterSpec 會被 validation 擋 → 保留完整 spec、執行中中止，避開 VSP / VCFMS"],
             ["Supervisor 網路", "VDS + Foundation LB (one-arm)", "NSX VPC 只支援 Avi；NSX segment 需 Edge + T0；FLB 內建於 vCenter，air-gap 免下載"],
             ["Registry", "外部 Harbor（單台 Linux）", "沒有 Software Depot；不在 Supervisor 內放 Harbor Supervisor Service"],
             ["Workload 網段", "192.168.130.0/24（VLAN 0）", "管理網與 workload 網必須不同 subnet；沿用 dcxd 既有 130.1 gateway，pfSense 不改"],
             ["addon 套件庫", "官方 override：新 AddonRepository + skip annotation", "builtin URL 寫死為公開網址；外部 registry 只能裝完後改指"],
             ["air-gap 驗證", "DNS 層級切斷（dnsmasq 無 upstream）", "pfSense 不可動；VKS 抓套件與 image 都用主機名稱"],
             ["nested vSAN", "silence vSAN 測試 nvmeonhcl", "HCL 黃燈會擋 vLCM 套用 Spherelet（僅限 lab）"]],
            [1.3, 2.6, 5.6], "這一頁是整份簡報的骨架：每個選擇背後都有被實測證明的理由。", font=11, rowh=0.52)
table_slide("Pick the Load Balancer That Needs Nothing Extra", "Supervisor 網路 × LB 相容矩陣",
            [["Supervisor 網路", "可用 LB", "額外需求", "本案"],
             ["VDS", "Foundation LB 或 Avi", "FLB：無（vCenter 內建，精靈 inline 部署）", "✔ 採用 FLB（Medium × 2, active/passive）"],
             ["NSX segment（T0/T1）", "NSX LB 或 Avi", "Edge cluster + T0 / T1", "—"],
             ["NSX VPC（bring-up 預設建立）", "只有 Avi", "Avi Controller OVA + 授權", "未使用（VPC 的 external IP block = 120.0/24，勿建 VPC）"]],
            [2.2, 2.0, 3.6, 3.7],
            "FLB 只能搭配 VDS 網路。雖然 bring-up 建了 NSX VPC，但 VPC Supervisor 只支援 Avi，所以我們選 VDS + FLB。"
            "精靈中 VDS 選項的相容性檢查，在 NSX prep 過的主機上也回 compatible。", font=12, rowh=0.62)
table_slide("Two Subnets, One L2: Address Plan", "IP 規劃（alan.lab）",
            [["用途", "位址", "用途", "位址"],
             ["vCenter / SDDC Manager / NSX", "120.3 / 120.4 / 120.11", "Workload 網（primary-wl）", "130.100 – .179，gw 130.1"],
             ["ESXi esx01–03", "120.61 – .63", "FLB VIP pool", "130.200 – .250"],
             ["Supervisor CP（5 連續）", "120.90 – .94", "Supervisor API VIP", "130.200（supervisor01.alan.lab）"],
             ["FLB 管理 IP", "120.95 – .97", "vks-test01 API VIP", "130.202"],
             ["Harbor / admin", "120.80 / 120.81", "VKS pods / services", "10.244.0.0/16 / 10.97.0.0/16"],
             ["bastion", "110.86（internet 側）", "Supervisor service CIDR", "10.96.0.0/23"],
             ["DNS（dnsmasq）", "120.5 / 110.85", "gateway", "120.1 / 130.1（pfSense dcxd）"]],
            [2.6, 2.4, 2.6, 3.2],
            "兩個網段在同一個 L2（VLAN 0），但 Supervisor 的管理網與 workload 網官方要求必須不同 subnet。"
            "VKS 預設 pod CIDR 192.168.0.0/16 會蓋掉 lab 網段，必須明確指定。", font=12, rowh=0.52)

# ---------------- section 2
section("Section - Lifecycle", "Platform: Bring-up to NSX, then Stop", "平台建置與 Supervisor 啟用",
        "第二章：用 VCF Installer 的 bring-up 取得 vCenter、vSAN、NSX，但在 VCF Management Platform 之前中止。")
diagram_slide("Get vCenter, vSAN and NSX — Skip VCF Management Services", "VCF bring-up 跑到 M4（NSX）即中止",
              "d3-bringup-stop.png",
              ["validation 強制要求 vspClusterSpec", "M1–M4 約 1 小時 45 分",
               "subtask #272 啟動時觸發中止", "installer 停、SDDC Manager 重啟 domainmanager",
               "結果只剩 vcsa / sddc / nsx01"],
              "為什麼不直接把 VSP 拿掉？因為 validation 會擋。所以保留完整 spec，用 watcher 在 M5 開始時中止。"
              "注意未開始的 subtask 狀態是 INITIALIZED，watcher 要把它當成閒置。")
table_slide("Enable the Supervisor with VDS and Foundation LB", "Supervisor 啟用精靈設定",
            [["步驟", "設定"],
             ["1 網路", "vSphere Distributed Switch (VDS)  ·  Load Balancer Type = Foundation Load Balancer"],
             ["2 位置", "Cluster deployment · supervisor01 · control plane HA（3 台）· vSphere Zone zone-m01"],
             ["3 Storage", "Control Plane / Ephemeral / Image Cache 皆為 vcf-m01-cl01 vSAN Storage Policy"],
             ["4 管理網", "Static · SDDC-DPortGroup-VM-Mgmt · 120.90 – 120.94 · gw 120.1 · DNS 120.5 · NTP 110.1"],
             ["5 Workload 網", "Static · pg-vks-workload → primary-wl · 130.100 – .179 · gw 130.1 · service CIDR 10.96.0.0/23"],
             ["6 Load Balancer", "flb-supervisor01 · One Arm · Medium · HA 2 · 管理網 flb-mgmt 120.95 – .96（後加 .97）· VIP 130.200 – .250"],
             ["7 進階", "Control plane Small（4 vCPU / 16 GB）· API DNS name supervisor01.alan.lab"]],
            [1.4, 10.1],
            "精靈操作重點：FLB 管理網要按 ASSIGN 另外定義且只能 2 個 IP；one-arm 的 Virtual Server 網路預設就是 primary workload 網。"
            "摘要頁不顯示 VIP、CP 大小、API DNS name，但值都有帶到。", font=12, rowh=0.55)
text_slide("Unblock Spherelet by Silencing the Nested vSAN HCL Check", "nested lab 地雷：vLCM 健康檢查擋住 Spherelet 8 小時",
           ["症狀：Supervisor RUNNING 但 k8s WARNING「No node is accepting vSphere Pods」，cci-ns Pending",
            "根因：Supervisor 以 vLCM ClusterApplySolution 安裝 Spherelet，FINALIZE 健康檢查 vSAN NOT_OK",
            "vSAN 測試「NVMe device is VMware certified」— nested 虛擬 NVMe 不在 VCG → 479 次失敗重試",
            "修法（僅限 lab）：SetVsanClusterSilentChecks(addSilentChecks=['nvmeonhcl'])，ID 要用短名",
            "結果：約 4 分鐘內 esx01–03 成為 agent node，cci-ns Running，Supervisor READY",
            "除錯入口：VCSA /var/log/vmware/vmware-updatemgr/vum-server/ 搜尋 'returned status: NOT_OK'"],
           "我一開始誤判成「VDS 不支援 vSphere Pod」，其實 9.1 VDS+FLB 的 vsphere_pods_enabled = True。黃燈看似無害，卻會擋住 vLCM。")

# ---------------- section 3
section("Section Header – Plum", "Air-Gap Supply Chain", "素材供應鏈：什麼東西、從哪裡、帶到哪裡",
        "第三章：所有素材由 bastion 下載，經 SHA256 驗證後帶進 air-gap 區。")
diagram_slide("Carry Every Artifact Across the Gap — with Checksums", "五類素材：來源 → bastion → admin → 目標",
              "d2-supply-chain.png",
              ["每一跳都用 SHA256SUMS 驗證", "harbor / admin 從不 apt、從不上網",
               "VCF CLI 9.1 只能從 Support Portal（公開 artifactory 只到 9.0.2）",
               "kubectl / kubectl-vsphere 直接從 Supervisor 下載",
               "Supervisor 無 Software Depot → /wcp/vcf-cli 固定回 503"],
              "Support Portal 路徑：My Downloads → VMware Cloud Foundation → 9.1.x → VCF Consumption CLI 與 VCF Consumption CLI Plugins。"
              "E:\\VKS 的 Supervisor zip 是 Supervisor 映像的 content library 包，不是 CLI。")
text_slide("Stand Up Harbor Offline and Make Everyone Trust It", "Harbor v2.15.2 離線安裝與 CA 信任",
           ["Docker .deb（5 個）與 Harbor offline installer 由 bastion 帶入，dpkg -i + install.sh",
            "自簽 CA「alan.lab Harbor Root CA」+ 伺服器憑證 SAN：harbor.alan.lab / harbor / 192.168.120.80",
            "信任點 1 — admin：/etc/docker/certs.d + 系統 trust store",
            "信任點 2 — Supervisor：container-image-registries API 登錄 Harbor（CA + 帳密）",
            "信任點 3 — VKS 節點：Cluster 變數 osConfiguration.trust.additionalTrustedCAs",
            "驗證：push library/airgap-test → admin 以 TLS 驗證 pull 成功；不裝 Trivy（DB 需外網）"],
           "三個信任點缺一不可。VKS 節點的信任靠 Cluster 變數，節點上 curl https://harbor.alan.lab 回 200 即為證明。")
text_slide("Kubernetes Releases from a Local Content Library", "VKr 匯入與相容性",
           ["來源 wp-content.broadcom.com/v2/latest（138 個 item）→ 每個 minor 最新 Photon 5",
            "v1.34.9 · v1.35.6 · v1.36.2，每個 item = vmdk + ovf + mf + cert",
            "govc library.import 匯入 vks-kr-local，item 名稱必須等於原始名稱",
            "設為 Supervisor default_kubernetes_service_content_library（非同步，約 20 秒生效）",
            "Namespace vks-ns01 綁定 content library 與 best-effort VM class",
            "kubectl get kr：三個版本皆 READY + COMPATIBLE"],
           "Supervisor 首頁提示的「Supervisor 映像 content library」是選配，用於 Supervisor 非同步升級，跟 VKr 的 content library 是兩回事。")

# ---------------- section 4
section("Section - VKS", "VKS Cluster & Add-on Repository", "VKS 叢集與 addon 套件庫",
        "第四章：建 VKS 叢集，並把 addon 套件庫改指內部 Harbor，最後在切斷外網下證明。")
code_slide("Provision VKS with kubectl — ClusterClass v3.7.0", "vks-test01：CAPI v1beta2 關鍵欄位",
           "apiVersion: cluster.x-k8s.io/v1beta2\nkind: Cluster\nmetadata: {name: vks-test01, namespace: vks-ns01}\nspec:\n"
           "  clusterNetwork:\n    pods:     {cidrBlocks: [\"10.244.0.0/16\"]}\n    services: {cidrBlocks: [\"10.97.0.0/16\"]}\n"
           "  topology:\n    classRef:            # v1beta1 的 topology.class\n      name: builtin-generic-v3.7.0\n"
           "      namespace: vmware-system-vks-public\n    version: v1.34.9+vmware.2-vkr.4\n    controlPlane: {replicas: 1}\n"
           "    workers:\n      machineDeployments:\n      - {class: node-pool, name: np1, replicas: 1}\n    variables:\n"
           "    - {name: vmClass, value: best-effort-medium}\n    - {name: storageClass, value: vcf-m01-cl01-vsan-storage-policy}\n"
           "    - name: osConfiguration\n      value:\n        trust:\n          additionalTrustedCAs:\n"
           "          - caCert: {content: <Harbor CA PEM>}",
           ["classRef 取代 v1beta1 的 class", "明確指定 CIDR，避開 192.168.x",
            "Harbor CA 經 osConfiguration.trust 注入節點", "先 kubectl apply --dry-run=server",
            "預設 Pod Security = restricted:latest"],
           "variable 定義要看 ClusterClass 的 status.variables[].definitions[]，不是 spec.variables。約 11 分鐘 Available，FLB 發 VIP 130.202。")
diagram_slide("Keep Add-ons Internal with the Documented Repository Override", "VKS 3.7 addon 套件庫：預設公開網址 → 改指 Harbor",
              "d4-addon-repo-override.png",
              ["builtin URL 寫死在 VKS bundle，無設定可改", "新 AddonRepository：只改 name + imageURL",
               "ARI 加 skip annotation、改 addonRepositoryRef", "workload 端由 carvel-repo 自動跟上",
               "VKS 升級後要再做一次"],
              "官方 techdocs「Update an Add-on Repository Install」即此做法。preferred-repository annotation 受 webhook 保護，不要改。"
              "沒有 VCFMS 時，這是外部 registry 唯一的內網路徑。")
table_slide("Proven Internet-Free: Negative Control and Recovery", "DNS air-gap 下的重做與驗證（2026-09-28）",
            [["步驟", "動作", "結果"],
             ["1 切斷外網", "dnsmasq 註解 server=8.8.8.8 / 1.1.1.1（pfSense 不動）", "公開名稱解析失敗；harbor / supervisor01 / vcsa 正常"],
             ["2 現狀驗證", "觀察 30 分鐘（pkgr 預設每 10 分同步）", "Supervisor 與 workload 皆於 T0 後多次從 Harbor 同步成功"],
             ["3 反向對照", "移除 skip annotation（先 dry-run）", "約 30 秒 ARI 退回 builtin → 兩端失敗：lookup projects.packages.broadcom.com；既有 addon 不受影響"],
             ["4 清理重建", "刪除舊 AR，由 builtin 重建乾淨 AR", "kapp 引用 0、公開網址 0；只保留 pkgr-type label 與 package-offerings"],
             ["5 改指", "skip=true + addonRepositoryRef → Harbor AR", "約 30 秒兩端恢復；下一輪 15:53Z 同步亦成功"],
             ["6 驗證", "controller log、狀態、image 來源", "'skipped updating AddonRepositoryRef'；0 個外部 image"]],
            [1.3, 4.3, 5.9],
            "反向對照證明兩件事：產品預設確實依賴外網；我們的 override 確實不依賴外網。每個修改都先 --dry-run=server。", font=11.5, rowh=0.62)
table_slide("Every Image Comes from Inside the Boundary", "最終驗收",
            [["驗收項目", "結果"],
             ["Supervisor", "k8s READY、無警告；3 CP + 3 ESXi agent；tkg 3.7.0 / velero / cci-ns 皆 CONFIGURED"],
             ["VKS cluster", "vks-test01 CP + worker Ready；AddonsReconciled=True；12 個 PackageInstall 全成功"],
             ["addon 套件庫", "Supervisor 與 workload PackageRepository 皆來自 harbor.alan.lab/vks/vks-addons"],
             ["Supervisor images", "localhost:5000 · docker-registry.kube-system.svc:5000 · vmware/* · etcd"],
             ["Workload images", "localhost:5000（VKr 內建）· harbor.alan.lab → 0 個外部 registry"],
             ["節點信任 / 拉取", "節點 curl https://harbor.alan.lab = 200；pod 從 Harbor 拉 226 MB image 成功"],
             ["工具", "VCF CLI 9.1.1 + 13 plugins：vcf context / cluster / addon 皆可用"]],
            [2.2, 9.3], "驗收時比對 registry 的集合，不比數量；要包含 initContainers。", font=12, rowh=0.55)

# ---------------- section 5
section("Section - Wrap Up", "Lessons Learned", "經驗、地雷與限制", "最後一章：這次踩到的雷，以及仍然存在的限制。")
table_slide("Pitfalls We Hit — and How to Avoid Them", "地雷總表",
            [["地雷", "避免方式"],
             ["bring-up spec 缺 vspClusterSpec 被擋", "保留完整 spec，M5 開始時中止；watcher 把 INITIALIZED 當閒置"],
             ["管理網與 workload 網同 subnet", "官方不支援 → workload 用另一個 subnet（本案 130.x）"],
             ["nested vSAN HCL 黃燈擋 vLCM", "silence nvmeonhcl（短 ID）；除錯看 vum-server log"],
             ["VKS 3.7 addon 預設公開網址", "Supervisor 端「成功」可能是偷用外網 → 以 image 來源與切網驗證為準"],
             ["dnsmasq --local-service 不回 routed subnet", "加 listen-address（參數仍會顯示，以實測為準）"],
             ["備份檔放在 /etc/dnsmasq.d/ 被當設定載入", "備份一律放設定目錄外"],
             ["VKS 預設 pod CIDR 192.168.0.0/16", "Cluster 明確指定 pods / services CIDR"],
             ["Pod Security 預設 restricted", "pod 需 runAsNonRoot、seccomp、drop ALL"],
             ["VCF CLI 從 Supervisor 下載 503", "沒有 Software Depot 就只能 Support Portal"]],
            [4.4, 7.1], "這些地雷都寫進狀態檔與記憶，重做時先看這張表。", font=11.5, rowh=0.46)
text_slide("Limitations and Next Steps", "限制與後續工作",
           ["builtin AddonRepository（公開網址）由 kapp 管理，會一直存在，但未被引用",
            "VKS 升級（如 add-on 3.7.0+20260723）會新增指向公開網址的 ARI → 升級後照同法改指",
            "DNS 層級切斷 ≠ 封包阻擋：直接用 IP 仍可出網；正式環境請搭配防火牆",
            "NSX VPC external IP block = 192.168.120.0/24 → 勿建 VPC 或先縮小範圍",
            "後續：用 vcf addon 安裝 cert-manager / contour；air-gap 升級 cluster 至 v1.35.6",
            "後續：以官方 PSO kit 產出正式安裝手冊（docx）"],
           "限制都已知且有對策；下一步是 addon 安裝與 cluster 升級的 air-gap 驗證。")
s = new("Big Statement - Plum")
ph(s, 12).text = "沒有 VCF Management Services，一樣能做完整 air-gap 的 VKS —— 用官方文件化的步驟，並以證據驗證。"
notes(s, "結論一句話。")

from pptx.oxml.ns import qn
def _runs(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _runs(sh.shapes)
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                yield from p.runs
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for p in c.text_frame.paragraphs:
                        yield from p.runs
for sl in list(prs.slides)[new_start - 1:]:
    for r in _runs(sl.shapes):
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea",):
            for old in rPr.findall(qn(tag)): rPr.remove(old)
        ea = rPr.makeelement(qn("a:ea"), {"typeface": "Microsoft JhengHei"})
        # schema order: ln, fill, effect, highlight, uLnTx, uLn, uFillTx, uFill, latin, ea, cs, sym, hlinkClick...
        latin = rPr.find(qn("a:latin"))
        if latin is not None: latin.addnext(ea)
        else:
            after = [rPr.find(qn(t)) for t in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")]
            after = [a for a in after if a is not None]
            if after: after[0].addprevious(ea)
            else: rPr.append(ea)

last_new = len(prs.slides)
prs.save(TMP)
order = [1] + list(range(new_start, last_new + 1)) + [28]
subprocess.run([sys.executable, os.path.join(HERE, "clone-slide.py"), "keep", TMP, "--slides", ",".join(map(str, order)), "--out", OUT], check=True)
# cover text
subprocess.run([sys.executable, os.path.join(HERE, "clone-slide.py"), "settext", OUT, "--slide", "1", "--out", OUT,
                "--title", "VKS Air-Gapped Deployment on VCF 9.1.1",
                "--subtitle", "不部署 VCF Management Services 的 VKS 離線安裝實作｜Supervisor (VDS + Foundation LB) + 外部 Harbor",
                "--set", "body:11=alan.lab 實驗環境", "--set", "body:12=Technical Runbook", "--set", "body:14=2026 年 9 月"], check=True)
print("saved", OUT, "slides:", len(order))
